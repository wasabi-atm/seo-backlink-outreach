"""Builds 14_video_seo_company_data.xlsx (user asked, 30 Sep): every checked company in one workbook, one tab
per video status, sorted by Ahrefs Domain Rating (highest first), company data only (no contacts,
campaign or note columns).

    python build_14_xlsx.py        (any Python 3.9+ with openpyxl)

Tabs:
  Video true, schema false   07 (126) plus the added 34 with a real video and no VideoObject
  Video true, schema true    08 plus the added 34 whose video has a VideoObject
  Video false                09 plus the added 34 with no real video (silent loops only)
The added 34 use the 30 Sep agent check (added34/video_result.csv), which replaces the older row where
a company was also in 08 or 09. Salesforce is left out: it blocks local renders, so its video is unknown.

Sources: Ahrefs connector (official API, batch-analysis) for DR and organic traffic; Similarweb through
Apify (radeance/similarweb-scraper for the 160, karamelo/similarweb-scraper-light for the other 860,
same numbers) for visits, bounce rate, average visit duration and pages per visit; Apollo free company
search for the staff band (Startup up to 50, Scaleup 51 to 500, Enterprise 501+).
"""
import csv, glob, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def rows(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


COLS = [("Company", 28), ("Domain", 26), ("Country", 9), ("Category", 34), ("Company size", 18),
        ("Video platform", 22), ("Schema found via", 30), ("Domain Rating", 13), ("Organic traffic", 15),
        ("Monthly visits", 15), ("Bounce rate", 12), ("Avg visit duration", 17), ("Pages per visit", 14)]
SIZE = {"1-50": "Startup (1-50)", "51-500": "Scaleup (51-500)", "501+": "Enterprise (501+)"}

# Ahrefs
dr = {r["webDomain"]: r for r in rows("metrics/ahrefs_api.csv") + rows("metrics/ahrefs_api_860.csv")}
# Similarweb: normalise both actors to visits, bounce (0-1), seconds, pages
sw = {}
for f in sorted(glob.glob("metrics/similarweb*.json"), key=lambda p: "test" not in p):
    for it in json.load(open(f)):
        d = (it.get("domain") or "").lower().removeprefix("www.")
        if not d:
            continue
        visits = it.get("totalVisits", it.get("monthlyVisits"))
        rec = dict(visits=visits, bounce=it.get("bounceRate"), secs=it.get("timeOnSite", it.get("avgVisitDurationSec")),
                   pages=it.get("pagesPerVisit"))
        if d not in sw or visits:
            sw[d] = rec
# Staff band
band = {r["webDomain"]: r["apolloBand"] for r in rows("metrics/apollo_bands.csv") + rows("metrics/apollo_bands_860.csv")}

added = {r["webDomain"]: r for r in rows("added34/targets.csv")}
agent = {r["webDomain"]: r for r in rows("added34/video_result.csv")}
tabs = {"Video true, schema false": [], "Video true, schema true": [], "Video false": []}
for f, tab in [("07_video_true_schema_false.csv", "Video true, schema false"),
               ("08_video_true_schema_true.csv", "Video true, schema true"),
               ("09_video_false_schema_false.csv", "Video false")]:
    for r in rows(f):
        if r["webDomain"] in added:
            continue  # the 30 Sep agent check replaces the older row
        tabs[tab].append(dict(company=r["companyName"], domain=r["webDomain"], country=r["country"],
                              category=r["serviceType"], platform=r["videoPlatform"] if tab != "Video false" else "",
                              schema=r["videoSchemaSource"] if tab == "Video true, schema true" else ""))
for d, a in agent.items():
    if a["hasVideo"] == "yes" and a["videoSchema"] == "FALSE":
        tab = "Video true, schema false"
    elif a["hasVideo"] == "yes" and a["videoSchema"] == "TRUE":
        tab = "Video true, schema true"
    elif a["hasVideo"] == "no":
        tab = "Video false"
    else:
        continue  # Salesforce: unknown
    t = added[d]
    tabs[tab].append(dict(company=t["companyName"], domain=d, country="", category=t["category"],
                          platform=a["videoPlatform"] if tab != "Video false" else "",
                          schema="VideoObject in the page" if tab == "Video true, schema true" else ""))
# country for the added 34 from the grouped file (Similarweb HQ)
hq = {r["webDomain"]: r["country"] for r in rows("12_video_seo_by_scale.csv")}

wb = Workbook()
wb.remove(wb.active)
head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F2937")
for name, data in tabs.items():
    for r in data:
        a, s = dr.get(r["domain"], {}), sw.get(r["domain"], {})
        has_sw = bool(s.get("visits"))
        r.update(country=r["country"] or hq.get(r["domain"], ""),
                 size=SIZE.get(band.get(r["domain"]), ""),
                 dr=float(a["domainRating"]) if a.get("domainRating") not in (None, "") else None,
                 traffic=int(float(a["organicTraffic"])) if a.get("organicTraffic") not in (None, "") else None,
                 visits=int(s["visits"]) if has_sw else None,
                 bounce=s.get("bounce") if has_sw else None,
                 secs=s.get("secs") if has_sw else None,
                 pages=round(s["pages"], 1) if has_sw and s.get("pages") is not None else None)
    data.sort(key=lambda r: (r["dr"] is None, -(r["dr"] or 0), -(r["traffic"] or 0)))
    ws = wb.create_sheet(name)
    for i, (title, width) in enumerate(COLS, 1):
        c = ws.cell(row=1, column=i, value=title)
        c.font, c.fill = head_font, head_fill
        ws.column_dimensions[get_column_letter(i)].width = width
    for n, r in enumerate(data, 2):
        vals = [r["company"], r["domain"], r["country"], r["category"], r["size"], r["platform"], r["schema"],
                r["dr"], r["traffic"], r["visits"], r["bounce"],
                r["secs"] / 86400 if r["secs"] is not None else None, r["pages"]]
        for i, v in enumerate(vals, 1):
            ws.cell(row=n, column=i, value=v)
        ws.cell(row=n, column=8).number_format = "0.#"
        ws.cell(row=n, column=9).number_format = "#,##0"
        ws.cell(row=n, column=10).number_format = "#,##0"
        ws.cell(row=n, column=11).number_format = "0.0%"
        ws.cell(row=n, column=12).number_format = "[m]:ss"
        ws.cell(row=n, column=13).number_format = "0.0"
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{len(data) + 1}"
    print(f"{name}: {len(data)} companies; no DR {sum(r['dr'] is None for r in data)}, "
          f"no Similarweb {sum(r['visits'] is None for r in data)}, no size {sum(not r['size'] for r in data)}")
wb.save("14_video_seo_company_data.xlsx")
print("saved 14_video_seo_company_data.xlsx")
